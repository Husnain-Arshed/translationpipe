import streamlit as st
import os
import io
import time
import markdown
from google import genai
from google.genai import types
from docx import Document
from docx.shared import Pt
from xhtml2pdf import pisa
from PIL import Image
from pdf2image import convert_from_bytes
from dotenv import load_dotenv


import streamlit as st

def check_password():
    """Restituisce True se l'utente ha inserito la password corretta."""
    def password_entered():
        # Controlla se la password inserita corrisponde a quella nei Secrets
        if st.session_state["pwd_input"] == st.secrets["APP_PASSWORD"]:
            st.session_state["password_correct"] = True
            del st.session_state["pwd_input"] # Cancella la password dalla memoria
        else:
            st.session_state["password_correct"] = False

    if st.session_state.get("password_correct", False):
        return True

    st.title("🔒 Accesso Riservato CAF")
    st.text_input("Password", type="password", on_change=password_entered, key="pwd_input")
    if "password_correct" in st.session_state:
        st.error("Password errata.")
    return False

if not check_password():
    st.stop()  # Blocca l'esecuzione del resto dell'app se non autorizzato

# ... qui inizia il resto del tuo codice originale ...

load_dotenv()



def optimize_image(img, max_dim=1600):
    """Ottimizza un oggetto PIL Image riducendone risoluzione e peso."""
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    if max(img.size) > max_dim:
        img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85, optimize=True)
    return buf.getvalue()

def optimize_image_bytes(image_bytes, max_dim=1600):
    """Wrapper per ottimizzare i byte diretti di un'immagine."""
    img = Image.open(io.BytesIO(image_bytes))
    return optimize_image(img, max_dim)

st.set_page_config(page_title="CAF AI Doc Translator", layout="wide", page_icon="⚖️")
st.title("⚖️ CAF & Patronati - Traduttore Documenti Ufficiali")

api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    st.warning("⚠️ Chiave API non trovata nelle variabili d'ambiente.")
    api_key = st.text_input("Inserisci GEMINI_API_KEY:", type="password")

if not api_key:
    st.info("👋 Inserisci la tua API Key di Google Gemini per iniziare.")
    st.stop()

client = genai.Client(api_key=api_key)

SYSTEM_PROMPT = """Sei un traduttore ufficiale e giurato esperto in atti amministrativi, anagrafici e fiscali per CAF e patronati italiani.

IL TUO COMPITO:
1. IDENTIFICAZIONE: Nella primissima riga, scrivi:
   **[Tipo documento: <tipo es. Certificato di Matrimonio, Nascita, FRC, Busta paga>]**
2. TRADUZIONE IN ITALIANO: Traduci il testo in modo fedele, accurato, professionale e senza alcuna omissione.

REGOLE CRITICHE TASSATIVE:
- NOMI E DATI ANAGRAFICI: Nomi, cognomi, date di nascita, numeri di protocollo, codici CNIC/NADRA, passaporti e toponimi devono rimanere ESATTAMENTE come nell'originale.
- STRUTTURA E TABELLE: Se nel documento originale ci sono tabelle o elenchi, riproducili fedelmente usando la sintassi Tabelle Markdown (| colonna 1 | colonna 2 |).
- TERMINI TECNICI/RELIGIOSI: Per termini specifici, usa l'equivalente funzionale in italiano e lascia il termine originale tra parentesi.
- TIMBRI E FIRME: Se vedi timbri o firme leggibili, segnalali tra parentesi quadre: es. `[Timbro: Registrato Union Council]`.
"""

def create_docx_from_markdown(md_text):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10.5)
    doc.add_paragraph("―" * 40)
    
    lines = md_text.split("\n")
    table_lines = []
    in_table = False

    def flush_table(t_lines):
        if not t_lines: return
        rows = [[cell.strip() for cell in r.split("|")[1:-1]] for r in t_lines if not set(r.strip()).issubset({'-', '|', ' '})]
        if not rows: return
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for r_idx, row in enumerate(rows):
            for c_idx, cell_value in enumerate(row):
                table.cell(r_idx, c_idx).text = cell_value

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False
            
            if stripped.startswith("### "): doc.add_heading(stripped.replace("### ", ""), level=3)
            elif stripped.startswith("## "): doc.add_heading(stripped.replace("## ", ""), level=2)
            elif stripped.startswith("# "): doc.add_heading(stripped.replace("# ", ""), level=1)
            elif stripped.startswith("- ") or stripped.startswith("* "): doc.add_paragraph(stripped[2:], style='List Bullet')
            elif stripped: doc.add_paragraph(stripped)
                
    if in_table:
        flush_table(table_lines)

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_io.seek(0)
    return docx_io

def create_pdf_from_markdown(md_text):
    html_content = markdown.markdown(md_text, extensions=['tables'])
    
    # Cerca il font Unicode esteso su macOS per supportare arabo/urdu
    font_paths = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf"
    ]
    
    font_path = ""
    for fp in font_paths:
        if os.path.exists(fp):
            font_path = fp
            break

    styled_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @font-face {{ font-family: 'UniversalFont'; src: url('{font_path}'); }}
            @page {{ size: a4 portrait; margin: 2cm; }}
            body {{ font-family: 'UniversalFont', Arial, sans-serif; font-size: 10pt; color: #333; line-height: 1.5; }}
            h1 {{ font-size: 16pt; color: #1a365d; border-bottom: 2px solid #2b6cb0; padding-bottom: 4px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 10px; margin-bottom: 10px; }}
            th, td {{ border: 1px solid #cbd5e0; padding: 6px 8px; text-align: left; font-size: 9pt; }}
            th {{ background-color: #edf2f7; font-weight: bold; }}
            .header-tag {{ background-color: #ebf8ff; border-left: 4px solid #3182ce; padding: 8px; margin-bottom: 15px; }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """
    pdf_io = io.BytesIO()
    # Usiamo styled_html direttamente senza specificare encoding
    pisa.CreatePDF(styled_html, dest=pdf_io)
    pdf_io.seek(0)
    return pdf_io

uploaded_file = st.file_uploader("Trascina qui il documento (PDF, Scansione, PNG, JPG, DOCX)", type=["pdf", "png", "jpg", "jpeg", "docx"])

if uploaded_file:
    file_bytes = uploaded_file.read()
    file_ext = uploaded_file.name.split(".")[-1].lower()
    
    col_input, col_action = st.columns([2, 1])
    with col_input:
        st.write(f"📁 **File selezionato:** `{uploaded_file.name}` ({len(file_bytes)/1024:.1f} KB)")

    with col_action:
        st.markdown("### Azioni")
        translate_btn = st.button("🚀 Avvia Traduzione", type="primary", use_container_width=True)

    if translate_btn:
        parts = []
        with st.spinner("Preparazione e ottimizzazione file..."):
            if file_ext == "pdf":
                # Converte il PDF in immagini a 150 DPI per ridurre il carico, poi ottimizza ciascuna pagina
                images = convert_from_bytes(file_bytes, dpi=150)
                for img in images:
                    opt_bytes = optimize_image(img, max_dim=1600)
                    parts.append(types.Part.from_bytes(data=opt_bytes, mime_type="image/jpeg"))
            elif file_ext in ["png", "jpg", "jpeg"]:
                opt_bytes = optimize_image_bytes(file_bytes, max_dim=1600)
                parts.append(types.Part.from_bytes(data=opt_bytes, mime_type="image/jpeg"))
            elif file_ext == "docx":
                doc_temp = Document(io.BytesIO(file_bytes))
                raw_text = "\n".join([p.text for p in doc_temp.paragraphs if p.text.strip()])
                parts.append(raw_text)

            parts.append("Esegui l'identificazione, l'OCR fedele e la traduzione completa in italiano rispettando le regole.")

        st.write("🔄 **Elaborazione in corso...**")
        output_container = st.empty()
        full_text = ""
        success = False
        max_retries = 3
        
        for attempt in range(max_retries):
            try:
                response_stream = client.models.generate_content_stream(
                    model="gemini-3.6-flash",
                    contents=parts,
                    config=types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT, temperature=0.1)
                )

                for chunk in response_stream:
                    if chunk.text:
                        full_text += chunk.text
                        output_container.markdown(full_text + "▌")
                        
                output_container.markdown(full_text)
                success = True
                break
            except Exception as err:
                if "503" in str(err) and attempt < max_retries - 1:
                    st.warning(f"Server occupato. Nuovo tentativo tra {2 * (attempt + 1)} secondi...")
                    time.sleep(2 * (attempt + 1))
                else:
                    st.error(f"Errore durante l'elaborazione: {err}")
                    break
        
        if success:
            st.session_state["raw_translation"] = full_text
            st.success(f"✅ Traduzione completata!")
            time.sleep(1)
            st.rerun()

if "raw_translation" in st.session_state:
    st.divider()
    st.subheader("✏️ Revisione e Validazione")
    edited_translation = st.text_area("Testo Tradotto (Markdown)", value=st.session_state["raw_translation"], height=450)

    col_docx, col_pdf = st.columns(2)
    with col_docx:
        st.download_button("📥 Scarica Word (.docx)", data=create_docx_from_markdown(edited_translation), file_name=f"TRADOTTO_{uploaded_file.name.rsplit('.', 1)[0]}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    with col_pdf:
        st.download_button("📄 Scarica PDF", data=create_pdf_from_markdown(edited_translation), file_name=f"TRADOTTO_{uploaded_file.name.rsplit('.', 1)[0]}.pdf", mime="application/pdf", use_container_width=True)